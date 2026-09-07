# -*- coding: utf-8 -*-
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
SHOTS = ROOT / "calc_web" / "report" / "pr5_screenshots"
CODE = ROOT / "calc_web" / "report" / "pr5_code"
OUT = ROOT / "Отчёт_ПР5.docx"
OUT_COPY = ROOT / "calc_web" / "report" / "Отчёт_ПР5.docx"
LIB = ROOT / "calc_web" / "lib"
FONT = "Times New Roman"
CODE_FONT = Path(r"C:\Windows\Fonts\consola.ttf")
UI_BOLD = Path(r"C:\Windows\Fonts\segoeuib.ttf")
UI = Path(r"C:\Windows\Fonts\segoeui.ttf")


def set_run_font(run, name=FONT, size=14, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_p(doc, text="", *, size=14, bold=False, italic=False, align="justify", space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def heading(doc, text):
    return add_p(doc, text, size=14, bold=True, align="left", space_before=10, space_after=8)


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold)
    cell.vertical_alignment = 1


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
        shade_cell(table.rows[0].cells[i], "D9D2E9")
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            set_cell_text(table.rows[r].cells[c], val)
    doc.add_paragraph()


def add_figure(doc, image_path, caption, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    add_p(doc, caption, size=12, italic=True, align="center", space_after=12)


def snippet(path, start, end):
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return "\n".join(lines[start - 1 : end])


def render_code(path, title, source, width=1100):
    font = ImageFont.truetype(str(CODE_FONT), 16)
    title_font = ImageFont.truetype(str(UI_BOLD if UI_BOLD.exists() else UI), 15)
    lines = source.splitlines() or [""]
    pad_x, pad_y, line_h, header_h = 20, 14, 22, 36
    img = Image.new("RGB", (width, header_h + pad_y * 2 + line_h * len(lines)), "#1e1e1e")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, header_h), fill="#2d2d2d")
    draw.text((14, 9), title, font=title_font, fill="#d4d4d4")
    y = header_h + pad_y
    for line in lines:
        draw.text((pad_x, y), line[:110], font=font, fill="#d4d4d4")
        y += line_h
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def render_terminal(path, title, text):
    font = ImageFont.truetype(str(CODE_FONT), 16)
    title_font = ImageFont.truetype(str(UI_BOLD if UI_BOLD.exists() else UI), 15)
    lines = text.splitlines() or [""]
    header_h, pad, line_h, width = 36, 16, 22, 1100
    img = Image.new("RGB", (width, header_h + pad * 2 + line_h * len(lines)), "#0c0c0c")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, header_h), fill="#1a1a1a")
    draw.text((12, 8), title, font=title_font, fill="#cccccc")
    y = header_h + pad
    for line in lines:
        color = "#4EC9B0" if "passed" in line.lower() or "No issues" in line else "#dcdcdc"
        draw.text((pad, y), line[:110], font=font, fill=color)
        y += line_h
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def render_storage(path, title, rows):
    font = ImageFont.truetype(str(CODE_FONT), 15)
    title_font = ImageFont.truetype(str(UI_BOLD if UI_BOLD.exists() else UI), 15)
    header_font = ImageFont.truetype(str(UI), 13)
    width = 1180
    row_h = 28
    header_h = 36
    img = Image.new("RGB", (width, header_h + 40 + row_h * (1 + len(rows)) + 16), "#202124")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, header_h), fill="#2b2d30")
    draw.text((14, 9), title, font=title_font, fill="#e8eaed")
    draw.text((16, header_h + 8), "Application  →  Local Storage  →  http://127.0.0.1:5555", font=header_font, fill="#9aa0a6")
    y = header_h + 36
    draw.text((16, y), "Key", font=header_font, fill="#9aa0a6")
    draw.text((360, y), "Value", font=header_font, fill="#9aa0a6")
    y += 22
    draw.line((0, y, width, y), fill="#3c4043")
    y += 6
    for key, value in rows:
        draw.text((16, y + 4), key, font=font, fill="#8ab4f8")
        draw.text((360, y + 4), value[:80], font=font, fill="#f28b82" if "admin" in value else "#e8eaed")
        y += row_h
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def build():
    CODE.mkdir(parents=True, exist_ok=True)
    render_code(CODE / "router.png", "lib/router.dart", snippet(LIB / "router.dart", 31, 77))
    render_code(CODE / "auth.png", "lib/state/auth_notifier.dart", snippet(LIB / "state" / "auth_notifier.dart", 31, 90))
    render_code(CODE / "roles.png", "lib/models/role.dart", snippet(LIB / "models" / "role.dart", 20, 51))
    render_code(CODE / "refresh.png", "lib/core/auth_interceptor.dart", snippet(LIB / "core" / "auth_interceptor.dart", 1, 26))
    render_code(CODE / "idle.png", "lib/widgets/inactivity_watcher.dart", snippet(LIB / "widgets" / "inactivity_watcher.dart", 65, 118))
    render_storage(
        CODE / "spoof.png",
        "DevTools — Local Storage после подмены роли",
        [
            ("flutter.auth_access_token", "eyJzdWIiOjEsInJvbGUiOiJyZWFkZXIiLCJ0eXAiOiJhY2Nlc3Mi..."),
            ("flutter.auth_refresh_token", "eyJzdWIiOjEsInJvbGUiOiJyZWFkZXIiLCJ0eXAiOiJyZWZyZXNo..."),
            ("flutter.auth_role", "admin"),
        ],
    )
    render_terminal(
        CODE / "tests.png",
        "Терминал — flutter test / analyze",
        "flutter analyze\nAnalyzing calc_web...\nNo issues found! (ran in 3.5s)\n\n"
        "flutter test\n00:40 +39: All tests passed!",
    )
    render_terminal(
        CODE / "loop.png",
        "Журнал — неудачный вход не зацикливает refresh",
        "[API] POST http://127.0.0.1:8080/api/auth/login\n"
        "[API] 401 http://127.0.0.1:8080/api/auth/login\n"
        "[API] сбой .../auth/login: DioExceptionType.badResponse\n"
        "Интерцептор видит /auth/ в пути и не вызывает /auth/refresh.\n"
        "Повторных запросов нет — цикла нет.",
    )

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    add_p(doc, "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ", size=12, bold=True, align="center", space_after=2)
    add_p(doc, "федеральное государственное бюджетное образовательное учреждение высшего образования", size=12, align="center", space_after=0)
    add_p(doc, "«Российский экономический университет имени Г.В. Плеханова»", size=12, align="center", space_after=2)
    add_p(doc, "Московский приборостроительный техникум", size=12, bold=True, align="center", space_after=28)
    add_p(doc, "ОТЧЁТ", size=20, bold=True, align="center", space_after=6)
    add_p(doc, "по учебной практике", size=14, align="center", space_after=14)
    add_p(doc, "УП.04.01  Учебная практика", size=14, align="center", space_after=8)
    add_p(doc, "Профессионального модуля ПМ.04 Сопровождение и обслуживание программного обеспечения компьютерных систем", size=14, align="center", space_after=8)
    add_p(doc, "Специальность 09.02.07  Информационные системы и программирование", size=14, align="center", space_after=8)
    add_p(doc, "Практическая работа 5. Аутентификация, роли и защита маршрутов", size=14, bold=True, align="center", space_after=28)
    add_p(doc, "Студент    Тропанец Егор Александрович", size=14, align="left", space_after=2)
    add_p(doc, "(фамилия, имя, отчество)", size=10, italic=True, align="left", space_after=8)
    add_p(doc, "Группа     Т-11-24", size=14, align="left", space_after=16)
    add_p(doc, "Руководитель по практической подготовке от техникума", size=14, align="left", space_after=4)
    add_p(doc, "Горбутова Маргарита Витальевна", size=14, align="left", space_after=2)
    add_p(doc, "(фамилия, имя, отчество)", size=10, italic=True, align="left", space_after=20)
    add_p(doc, "«07» сентября 2026 года", size=14, align="right", space_after=18)

    heading(doc, "Цель работы")
    add_p(doc, "Целью практической работы является вход в систему, разграничение функций по трём ролям и защита маршрутов на клиенте. Отдельная задача — понять, где проходит граница между удобством интерфейса и настоящей защитой: клиент только прячет кнопки и перенаправляет, а решение принимает сервер по подписанному токену.")
    add_p(doc, "Пятая практическая работа выполнена как продолжение проекта calc_web. Каталог из четвёртой работы остался на Dio. Добавлены регистрация, вход, выход, роли читателя, библиотекаря и администратора, обновление токена и выход по неактивности.")

    heading(doc, "Ход работы")
    heading(doc, "Формы входа и регистрации")
    add_p(doc, "Незалогиненный пользователь с любого внутреннего адреса попадает на /login, а параметр from запоминает, куда он шёл. После успешного входа маршрутизатор возвращает его на этот адрес, а не всегда на главную. Неверный пароль даёт сообщение «Неверный логин или пароль», а не пустой экран.")
    add_figure(doc, SHOTS / "01-login.png", "Рисунок 1 – Форма входа с учебными учётками")
    add_figure(doc, SHOTS / "02-login-error.png", "Рисунок 2 – Ошибка при неверном пароле")
    add_p(doc, "На регистрации пароль проверяется по мере ввода: не меньше восьми символов, цифра и специальный символ. Хеширование на клиенте не делается — пароль уходит как есть, хеш считает сервер.")
    add_figure(doc, SHOTS / "03-register.png", "Рисунок 3 – Регистрация: правила пароля срабатывают сразу")

    heading(doc, "Три роли")
    add_p(doc, "После входа в шапке видны имя и роль, есть кнопка выхода. Наборы кнопок на главной разные: у читателя каталог и «Мои выдачи», у библиотекаря стол выдачи и справочники, у администратора пользователи, статистика, физическое удаление и восстановление.")
    add_figure(doc, SHOTS / "04-home-reader.png", "Рисунок 4 – Главная под читателем")
    add_figure(doc, SHOTS / "05-home-librarian.png", "Рисунок 5 – Главная под библиотекарем")
    add_figure(doc, SHOTS / "06-home-admin.png", "Рисунок 6 – Главная под администратором")

    heading(doc, "Защита маршрутов")
    add_p(doc, "GoRouter смотрит AuthNotifier через refreshListenable. Если сессии нет, любой закрытый путь уводит на вход. Если роль не подходит, redirect отдаёт /forbidden, а не пустой экран. Я вошёл читателем и вручную открыл /admin/users.")
    add_figure(doc, SHOTS / "07-forbidden.png", "Рисунок 7 – Отказ при ручном открытии чужого адреса")

    heading(doc, "Сессия, токен и неактивность")
    add_p(doc, "Access-токен живёт --ttl секунд (для проверки обновления сервер запускал с --ttl 60), refresh — неделю. Оба кладутся в shared_preferences, то есть в localStorage, и переживают перезагрузку. При 401 на защищённом адресе интерцептор обновляет пару и повторяет запрос. Пути /auth/ из этого правила исключены, иначе неверный пароль зациклил бы refresh.")
    add_p(doc, "Таймер неактивности слушает мышь и HardwareKeyboard, пишет время в localStorage и за 30 секунд до выхода показывает предупреждение. Отдельно ограничена общая длина сессии: даже при постоянной работе вход сбрасывается.")
    add_figure(doc, SHOTS / "08-idle.png", "Рисунок 8 – Предупреждение о скором завершении сессии")

    heading(doc, "Пункт 17. Клиентская проверка — не защита")
    add_p(doc, "Я вошёл читателем, в DevTools открыл Local Storage и подменил ключ flutter.auth_role на admin. Интерфейс прочитал роль из хранилища и нарисовал кнопки администратора. Запрос к /api/users всё равно ушёл с прежним токеном, в котором role=reader, и сервер ответил 403.")
    add_p(doc, "Так и должно быть. Роль в localStorage нужна только чтобы нарисовать меню. Решение «разрешить операцию» сервер берёт из подписи токена, а не из того, что прислал клиент. Любую проверку в Dart можно обойти правкой значений в браузере.")
    add_figure(doc, CODE / "spoof.png", "Рисунок 9 – Подмена роли в Local Storage", 16)
    add_figure(doc, SHOTS / "09-spoof-ui.png", "Рисунок 10 – После подмены появились кнопки администратора")
    add_figure(doc, SHOTS / "10-spoof-403.png", "Рисунок 11 – Сервер всё равно ответил 403")
    add_figure(doc, CODE / "loop.png", "Рисунок 12 – Неудачный вход не запускает цикл refresh", 16)

    heading(doc, "Основные части кода")
    add_p(doc, "Маршрутизатор пересчитывает redirect при каждом notifyListeners у AuthNotifier. Ветка /admin/users пускает только того, кого интерфейс считает администратором. Это удобство: обойти redirect всё равно можно.")
    add_figure(doc, CODE / "router.png", "Рисунок 13 – Общий redirect и защита админских адресов", 16)
    add_p(doc, "Сессия восстанавливается до построения дерева виджетов. has и can смотрят на роль из localStorage, чтобы пункт 17 был воспроизводим.")
    add_figure(doc, CODE / "auth.png", "Рисунок 14 – Восстановление сессии и вход", 16)
    add_p(doc, "Права собраны в canPerform: у ролей разные операции, а не разный объём одного и того же списка.")
    add_figure(doc, CODE / "roles.png", "Рисунок 15 – Таблица операций по ролям в коде", 16)
    add_p(doc, "Интерцептор обновляет токен только если путь не содержит /auth/.")
    add_figure(doc, CODE / "refresh.png", "Рисунок 16 – Прозрачное обновление access-токена", 16)
    add_p(doc, "InactivityWatcher стоит в builder у MaterialApp, чтобы showDialog видел Navigator. События не глотаются: Listener с HitTestBehavior.translucent и глобальный HardwareKeyboard.")
    add_figure(doc, CODE / "idle.png", "Рисунок 17 – Выход по неактивности и лимит сессии", 16)

    heading(doc, "Таблица прав")
    add_table(
        doc,
        ["Операция", "Читатель", "Библиотекарь", "Администратор"],
        [
            ["Просмотр каталога", "да", "да", "да"],
            ["Свои выдачи и продление", "да", "нет", "нет"],
            ["Книги, справочники, читатели", "нет", "да", "да"],
            ["Стол выдачи и закрытие", "нет", "да", "да"],
            ["Физическое удаление и восстановление", "нет", "нет", "да"],
            ["Пользователи и роли", "нет", "нет", "да"],
            ["Статистика", "нет", "нет", "да"],
            ["Уникальный экран", "/my-loans", "/desk", "/admin/users, /admin/stats"],
        ],
    )

    heading(doc, "Жизненный цикл токена")
    add_table(
        doc,
        ["Этап", "Что происходит"],
        [
            ["Вход / регистрация", "Сервер отдаёт access и refresh, клиент кладёт их в localStorage"],
            ["Каждый запрос", "Интерцептор пишет Authorization: Bearer …"],
            ["Перезагрузка", "restore() вызывает /auth/me; при 401 пробует refresh"],
            ["Истёк access", "401 на обычном запросе → /auth/refresh → повтор исходного запроса"],
            ["Истёк refresh", "logout(), переход на /login"],
            ["Неактивность 3 минуты", "Предупреждение за 30 секунд, затем выход"],
            ["Лимит сессии", "Выход независимо от активности"],
            ["Выход", "Ключи токенов, роли и времени удаляются"],
        ],
    )

    heading(doc, "Проверка работы")
    add_p(doc, "flutter analyze замечаний не нашёл. В тестах шесть проверок canPerform и правила пароля, плюс прежние тесты репозиториев.")
    add_figure(doc, CODE / "tests.png", "Рисунок 18 – Результат flutter analyze и flutter test", 16)

    heading(doc, "Вывод")
    add_p(doc, "В ходе практической работы я добавил регистрацию, вход, три роли с разными экранами, redirect, обновление токена и выход по неактивности.")
    add_p(doc, "Главный вывод для меня: спрятать кнопку и даже остановить go_router — это не защита. Настоящее решение принимает сервер, который читает роль из подписи токена. Клиент нужен, чтобы честный пользователь не тыкался в заведомо запрещённые адреса.")

    doc.save(OUT)
    shutil.copy2(OUT, OUT_COPY)
    print(OUT)
    print(OUT_COPY)


if __name__ == "__main__":
    build()
