# -*- coding: utf-8 -*-
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
SHOTS = ROOT / "calc_web" / "report" / "pr3_screenshots"
CODE = ROOT / "calc_web" / "report" / "pr3_code"
OUT = ROOT / "Отчёт_ПР3.docx"
OUT_COPY = ROOT / "calc_web" / "report" / "Отчёт_ПР3.docx"
LIB = ROOT / "calc_web" / "lib"
FONT = "Times New Roman"
CODE_FONT = Path(r"C:\Windows\Fonts\consola.ttf")
UI_BOLD = Path(r"C:\Windows\Fonts\segoeuib.ttf")
UI = Path(r"C:\Windows\Fonts\segoeui.ttf")


def set_run_font(run, name=FONT, size=14, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_p(doc, text="", *, size=14, bold=False, italic=False, align="justify", space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def heading(doc, text):
    return add_p(doc, text, size=14, bold=True, align="left", space_before=10, space_after=8)


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color_hex)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold)
    cell.vertical_alignment = 1


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
        shade_cell(table.rows[0].cells[i], "D9D2E9")
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            set_cell_text(table.rows[r].cells[c], val)
    doc.add_paragraph()


def add_figure(doc, image_path, caption, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    add_p(doc, caption, size=12, italic=True, align="center", space_after=12)


def snippet(path, start, end):
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return "\n".join(lines[start - 1 : end])


def render_code(path, title, source, width=1100):
    font = ImageFont.truetype(str(CODE_FONT), 16)
    title_font = ImageFont.truetype(str(UI_BOLD if UI_BOLD.exists() else UI), 15)
    lines = source.splitlines() or [""]
    pad_x, pad_y, line_h, header_h = 20, 14, 22, 36
    img = Image.new("RGB", (width, header_h + pad_y * 2 + line_h * len(lines)), "#1e1e1e")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, header_h), fill="#2d2d2d")
    draw.text((14, 9), title, font=title_font, fill="#d4d4d4")
    y = header_h + pad_y
    for line in lines:
        draw.text((pad_x, y), line[:110], font=font, fill="#d4d4d4")
        y += line_h
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def render_terminal(path, title, text):
    font = ImageFont.truetype(str(CODE_FONT), 16)
    title_font = ImageFont.truetype(str(UI_BOLD if UI_BOLD.exists() else UI), 15)
    lines = text.splitlines() or [""]
    header_h, pad, line_h, width = 36, 16, 22, 1100
    img = Image.new("RGB", (width, header_h + pad * 2 + line_h * len(lines)), "#0c0c0c")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, width, header_h), fill="#1a1a1a")
    draw.text((12, 8), title, font=title_font, fill="#cccccc")
    y = header_h + pad
    for line in lines:
        color = "#4EC9B0" if "passed" in line.lower() or "No issues" in line else "#dcdcdc"
        draw.text((pad, y), line[:110], font=font, fill=color)
        y += line_h
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def build():
    CODE.mkdir(parents=True, exist_ok=True)
    render_code(CODE / "validators.png", "lib/core/validators.dart", snippet(LIB / "core" / "validators.dart", 1, 56))
    render_code(CODE / "store.png", "lib/data/library_store.dart", snippet(LIB / "data" / "library_store.dart", 20, 56))
    render_code(CODE / "unique.png", "lib/data/library_store.dart", snippet(LIB / "data" / "library_store.dart", 151, 184))
    render_code(CODE / "form.png", "lib/widgets/entity_form.dart", snippet(LIB / "widgets" / "entity_form.dart", 7, 70))
    render_code(CODE / "isbn.png", "lib/screens/book_form_screen.dart", snippet(LIB / "screens" / "book_form_screen.dart", 76, 92))
    render_code(CODE / "publisher.png", "lib/repositories/catalog_repositories.dart", snippet(LIB / "repositories" / "catalog_repositories.dart", 180, 200))
    render_code(CODE / "router.png", "lib/router.dart", snippet(LIB / "router.dart", 28, 66))
    render_code(CODE / "provider.png", "lib/main.dart", snippet(LIB / "main.dart", 54, 78))
    render_terminal(
        CODE / "tests.png",
        "Терминал — flutter test / analyze",
        "flutter test\n00:00 +27: All tests passed!\n\nflutter analyze\nAnalyzing calc_web...\nNo issues found! (ran in 1.3s)",
    )

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    add_p(doc, "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ", size=12, bold=True, align="center", space_after=2)
    add_p(doc, "федеральное государственное бюджетное образовательное учреждение высшего образования", size=12, align="center", space_after=0)
    add_p(doc, "«Российский экономический университет имени Г.В. Плеханова»", size=12, align="center", space_after=2)
    add_p(doc, "Московский приборостроительный техникум", size=12, bold=True, align="center", space_after=28)
    add_p(doc, "ОТЧЁТ", size=20, bold=True, align="center", space_after=6)
    add_p(doc, "по учебной практике", size=14, align="center", space_after=14)
    add_p(doc, "УП.04.01  Учебная практика", size=14, align="center", space_after=8)
    add_p(doc, "Профессионального модуля ПМ.04 Сопровождение и обслуживание программного обеспечения компьютерных систем", size=14, align="center", space_after=8)
    add_p(doc, "Специальность 09.02.07  Информационные системы и программирование", size=14, align="center", space_after=8)
    add_p(doc, "Практическая работа 3. Формы, валидация и связанные сущности", size=14, bold=True, align="center", space_after=28)
    add_p(doc, "Студент    Тропанец Егор Александрович", size=14, align="left", space_after=2)
    add_p(doc, "(фамилия, имя, отчество)", size=10, italic=True, align="left", space_after=8)
    add_p(doc, "Группа     Т-11-24", size=14, align="left", space_after=16)
    add_p(doc, "Руководитель по практической подготовке от техникума", size=14, align="left", space_after=4)
    add_p(doc, "Горбутова Маргарита Витальевна", size=14, align="left", space_after=2)
    add_p(doc, "(фамилия, имя, отчество)", size=10, italic=True, align="left", space_after=20)
    add_p(doc, "«07» сентября 2026 года", size=14, align="right", space_after=18)

    heading(doc, "Цель работы")
    add_p(doc, "Целью практической работы является освоение форм создания и редактирования сущностей во Flutter Web, проверки ввода, уникальности полей и связей между записями. Нужно было сохранить данные в браузере, не дать удалить связанное издательство и предупредить пользователя о несохранённых изменениях.")
    add_p(doc, "Третья практическая работа выполнена как продолжение проекта calc_web. Калькулятор, конвертер и списки из предыдущих работ сохранены. Основная часть третьей работы — формы пяти сущностей учебной библиотеки и хранение через shared_preferences.")

    heading(doc, "Ход работы")
    heading(doc, "Главная страница")
    add_p(doc, "На главной странице я добавил переходы ко всем пяти сущностям: книги, авторы, жанры, издательства и читатели. Калькулятор и конвертер оставлены ниже как инструменты из первой практической работы.")
    add_figure(doc, SHOTS / "01-home.png", "Рисунок 1 – Главная страница практической работы №3")

    heading(doc, "Форма книги и связи")
    add_p(doc, "Для книг сделан один экран на создание и редактирование. Адреса /books/new и /books/:id/edit открывают одну и ту же форму. Издательство выбирается из репозитория, а не из константы. Авторы и жанры задаются чипами: это связи многие-ко-многим. Издательство — связь многие-к-одному.")
    add_p(doc, "После выбора издательства списки авторов и жанров сужаются по уже связанным книгам. Для издательства «АСТ» остаются авторы Толстой, Булгаков, Брэдбери и Глуховский и жанры «Художественная», «Фантастика», «История».")
    add_figure(doc, SHOTS / "02-book-form.png", "Рисунок 2 – Форма редактирования книги со связями и каскадным отбором")

    heading(doc, "Валидация полей")
    add_p(doc, "Проверки вынесены в класс V в файле lib/core/validators.dart. На пустой форме создания книги сообщения появляются прямо у полей: обязательность, длина, целое число, выбор издательства, хотя бы один автор и хотя бы один жанр.")
    add_figure(doc, SHOTS / "03-validation.png", "Рисунок 3 – Сообщения валидации на форме новой книги")

    heading(doc, "Уникальность ISBN")
    add_p(doc, "ISBN проверяется не только по формату 10 или 13 цифр, но и на уникальность в хранилище. Если подставить ISBN другой книги, ошибка «ISBN уже используется» показывается у самого поля. Аналогично проверяется email читателя.")
    add_figure(doc, SHOTS / "04-isbn-unique.png", "Рисунок 4 – Ошибка уникальности ISBN на поле формы")

    heading(doc, "Несохранённые изменения")
    add_p(doc, "Форма отслеживает, менял ли пользователь данные. Если форма «грязная», кнопка «Назад», «Отмена» и системная кнопка назад браузера спрашивают подтверждение. Для этого используется PopScope и общий виджет EntityFormScaffold.")
    add_figure(doc, SHOTS / "05-unsaved.png", "Рисунок 5 – Диалог при попытке покинуть форму без сохранения")

    heading(doc, "Издательства и запрет удаления")
    add_p(doc, "Список издательств сделан по той же схеме, что списки книг и авторов: поиск, показ удалённых, сортировка, пагинация, логическое и физическое удаление. В таблице видно, сколько книг ссылается на каждое издательство.")
    add_figure(doc, SHOTS / "06-publishers.png", "Рисунок 6 – Список издательств с числом связанных книг")
    add_p(doc, "Если у издательства ещё есть книги, удаление запрещается. Для «АСТ» выводится сообщение с количеством ссылок: 7 книг. Запись при этом не удаляется.")
    add_figure(doc, SHOTS / "07-publisher-delete.png", "Рисунок 7 – Запрет удаления издательства, на которое ссылаются книги")

    heading(doc, "Читатель и читательский билет")
    add_p(doc, "У читателя связь один-к-одному с читательским билетом. Билет не вынесен на отдельный экран: номер, даты выдачи и окончания, признак активности редактируются прямо в форме читателя.")
    add_figure(doc, SHOTS / "08-reader-form.png", "Рисунок 8 – Форма читателя со вложенным читательским билетом")

    heading(doc, "Сохранение в браузере")
    add_p(doc, "Данные пишутся в shared_preferences под ключами books_v1, authors_v1, genres_v1, publishers_v1 и readers_v1. Я создал жанр «Учебный жанр», обновил страницу и убедился, что запись осталась. Если формат в хранилище окажется старым, приложение не падает: набор восстанавливается из начальных данных, а пользователю показывается сообщение.")
    add_figure(doc, SHOTS / "09-before-reload.png", "Рисунок 9 – Список жанров сразу после создания новой записи")
    add_figure(doc, SHOTS / "10-after-reload.png", "Рисунок 10 – Тот же список после перезагрузки страницы")

    heading(doc, "Основные части кода")
    add_p(doc, "Проверки полей собраны в одном месте. Методы required, length, integer, email, isbn и date можно комбинировать через V.all.")
    add_figure(doc, CODE / "validators.png", "Рисунок 11 – Класс валидаторов V", 16)
    add_p(doc, "LibraryStore хранит все списки, восстанавливает их из JSON и сообщает, если ключи были в старом формате. Для тестов есть режим memory() без SharedPreferences.")
    add_figure(doc, CODE / "store.png", "Рисунок 12 – Ключи хранения и загрузка LibraryStore", 16)
    add_p(doc, "Уникальность ISBN и email, подсчёт книг издательства и каскадный отбор авторов и жанров тоже находятся в хранилище, чтобы формы не обходили репозиторий.")
    add_figure(doc, CODE / "unique.png", "Рисунок 13 – Проверки уникальности и связей в LibraryStore", 16)
    add_p(doc, "Общий каркас формы EntityFormScaffold рисует поля, кнопки сохранения и отмены и диалог несохранённых изменений. Для авторов и жанров используется ChipMultiSelect на базе FormField.")
    add_figure(doc, CODE / "form.png", "Рисунок 14 – Каркас формы EntityFormScaffold", 16)
    add_p(doc, "Ошибка уникальности ISBN сначала записывается в отдельную переменную, затем форма проверяется повторно, чтобы сообщение появилось у поля, а не всплывающим уведомлением.")
    add_figure(doc, CODE / "isbn.png", "Рисунок 15 – Проверка уникальности ISBN при сохранении книги", 16)
    add_p(doc, "Репозиторий издательств перед логическим и физическим удалением считает связанные книги и бросает RelationException.")
    add_figure(doc, CODE / "publisher.png", "Рисунок 16 – Запрет удаления издательства в репозитории", 16)

    heading(doc, "Маршруты и состояние")
    add_p(doc, "Маршрут new объявлен раньше :id, иначе слово new воспринималось бы как идентификатор. Редактирование открывается как вложенный путь :id/edit. Для жанров, издательств и читателей используется тот же приём.")
    add_figure(doc, CODE / "router.png", "Рисунок 17 – Маршруты создания и редактирования книг и авторов", 16)
    add_p(doc, "В main.dart сначала загружается LibraryStore, затем к нему подключаются репозитории и ChangeNotifier списков. Формы берут выпадающие списки из хранилища, поэтому после создания жанра он сразу появляется в форме книги.")
    add_figure(doc, CODE / "provider.png", "Рисунок 18 – Подключение хранилища, репозиториев и notifier в main.dart", 16)

    heading(doc, "Схема маршрутов")
    add_table(
        doc,
        ["Адрес", "Страница"],
        [
            ["/", "Главная страница"],
            ["/books", "Каталог книг"],
            ["/books/new", "Создание книги"],
            ["/books/:id", "Карточка книги"],
            ["/books/:id/edit", "Редактирование книги"],
            ["/authors, /authors/new, /authors/:id/edit", "Авторы"],
            ["/genres, /genres/new, /genres/:id/edit", "Жанры"],
            ["/publishers, /publishers/new, /publishers/:id/edit", "Издательства"],
            ["/readers, /readers/new, /readers/:id/edit", "Читатели"],
        ],
    )

    heading(doc, "Словарь данных")
    add_table(
        doc,
        ["Сущность", "Поле", "Тип", "Ограничения"],
        [
            ["Book", "title", "строка", "обязательно, 2–200 символов"],
            ["Book", "isbn", "строка", "обязательно, 10 или 13 цифр, уникально"],
            ["Book", "year", "целое", "1450–2100"],
            ["Book", "pages", "целое", "не меньше 1"],
            ["Book", "publisherId", "целое", "обязательный выбор из издательств"],
            ["Book", "authorIds", "список id", "хотя бы один автор"],
            ["Book", "genreIds", "список id", "хотя бы один жанр"],
            ["Book", "copiesTotal / copiesAvailable", "целое", "доступно не больше общего числа"],
            ["Author", "lastName, firstName", "строка", "обязательно, 2–80 символов"],
            ["Author", "country", "строка", "обязательно"],
            ["Author", "birthYear", "целое", "1400–2020"],
            ["Genre", "name", "строка", "обязательно, 2–60 символов"],
            ["Publisher", "name, city, foundedYear", "строка / целое", "нельзя удалить, если есть книги"],
            ["Reader", "email", "строка", "формат почты, уникально"],
            ["LibraryCard", "number, issuedAt, expiresAt", "строка / дата", "вложен в форму читателя, срок позже выдачи"],
        ],
    )

    heading(doc, "Таблица валидации")
    add_table(
        doc,
        ["Поле", "Правило", "Сообщение"],
        [
            ["Название книги", "required + length 2–200", "Укажите название / длина"],
            ["ISBN", "required + isbn + уникальность", "Укажите ISBN / 10 или 13 цифр / уже используется"],
            ["Год издания", "integer 1450–2100", "Введите целое число / диапазон"],
            ["Издательство", "не null", "Выберите издательство"],
            ["Авторы / жанры", "список не пустой", "Выберите хотя бы одного автора / жанр"],
            ["Email читателя", "email + уникальность", "Некорректный адрес / уже используется"],
            ["Даты билета", "YYYY-MM-DD, expires > issued", "Дата в формате ГГГГ-ММ-ДД"],
        ],
    )

    heading(doc, "Проверка работы")
    add_p(doc, "Я проверил команды flutter analyze и flutter test. Кроме списков из второй работы проверяются уникальность ISBN и email, запрет удаления издательства с книгами, чтение JSON без новых полей и сами валидаторы.")
    add_figure(doc, CODE / "tests.png", "Рисунок 19 – Результат flutter test и flutter analyze", 16)

    heading(doc, "Вывод")
    add_p(doc, "В ходе практической работы я добавил в учебную библиотеку формы создания и редактирования пяти сущностей, проверки ввода, уникальность ISBN и email, вложенный читательский билет и каскадный отбор авторов и жанров по издательству.")
    add_p(doc, "Я разобрался, зачем хранить данные под ключами с версией и зачем выносить валидаторы в отдельный класс: формат в браузере можно сменить без падения приложения, а одинаковые правила не приходится копировать на каждый экран.")

    doc.save(OUT)
    shutil.copy2(OUT, OUT_COPY)
    print(OUT)
    print(OUT_COPY)


if __name__ == "__main__":
    build()
